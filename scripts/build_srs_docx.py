"""Build the polished OSIEL developer SRS from its authoritative Markdown."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "17233F"
MUTED = "667085"
LIGHT = "F2F4F7"
VIOLET = "6657D9"
PALE_VIOLET = "F2F0FF"
GREEN = "16866B"
PALE_GREEN = "EAF8F4"
AMBER = "9A681C"
PALE_AMBER = "FFF5DE"
RED = "A94050"
PALE_RED = "FDECEF"
WHITE = "FFFFFF"
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_widths(table, widths_in: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    dxa_widths = [int(round(width * 1440)) for width in widths_in]
    # Make Word geometry exact even when decimal-inch patterns would round to
    # 9359/9361 DXA. The final column absorbs the one-DXA correction.
    dxa_widths[-1] += CONTENT_WIDTH_DXA - sum(dxa_widths)
    for width in dxa_widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = dxa_widths[index]
            cell.width = Inches(widths_in[index])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_font(run, size=11, color="222222", bold=None, italic=None, name="Calibri") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_runs(paragraph, text: str, size=11, color="222222", bold=False, italic=False) -> None:
    parts = re.split(r"(\*\*.*?\*\*|`[^`]+`|https?://\S+)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, size=size, color=color, bold=True, italic=italic)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, size=max(8, size - 1), color=DARK_BLUE, bold=bold, italic=italic, name="Consolas")
        else:
            run = paragraph.add_run(part)
            set_font(run, size=size, color=color, bold=bold, italic=italic)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_font(run, size=8, color=MUTED)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    properties = doc.core_properties
    properties.title = "Oncogon AI OSIEL — Developer-Grade Software Requirements Specification"
    properties.subject = "Scientific intelligence, governed Professor, 3D molecular visualization, model laboratory and experimental learning engine specification"
    properties.author = "Oncogon AI"
    properties.keywords = "OSIEL, oncology, cheminformatics, compound prioritization, SRS"
    properties.comments = "Research use only. Generated from the authoritative repository Markdown."

    header = section.header
    header_table = header.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_widths(header_table, [3.9, 2.6])
    header_table.cell(0, 0).text = "ONCOGON AI  ·  OSIEL"
    header_table.cell(0, 1).text = "DEVELOPER SRS  ·  v1.2"
    for index, cell in enumerate(header_table.rows[0].cells):
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.RIGHT
        for run in cell.paragraphs[0].runs:
            set_font(run, size=8, color=VIOLET if index == 0 else MUTED, bold=True)
    header_table._tbl.tblPr.remove(header_table._tbl.tblPr.find(qn("w:tblBorders"))) if header_table._tbl.tblPr.find(qn("w:tblBorders")) is not None else None

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer_p.add_run("RESEARCH USE ONLY  ·  ")
    set_font(run, size=8, color=MUTED, bold=True)
    add_page_field(footer_p)


def add_cover(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.LEFT
    kicker.paragraph_format.space_after = Pt(8)
    add_runs(kicker, "ONCOGON AI  /  TECHNICAL SPECIFICATION", size=10, color=VIOLET, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(8)
    add_runs(title, "OSIEL", size=34, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(4)
    add_runs(subtitle, "Scientific Intelligence & Experimental Learning Engine", size=18, color=DARK_BLUE, bold=True)

    descriptor = doc.add_paragraph()
    descriptor.paragraph_format.space_after = Pt(22)
    add_runs(descriptor, "Developer-Grade Software Requirements Specification", size=14, color=MUTED)

    callout = doc.add_table(rows=1, cols=1)
    set_table_widths(callout, [6.5])
    set_cell_shading(callout.cell(0, 0), PALE_VIOLET)
    p = callout.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_runs(p, "67-point core scope + implementation addenda  ·  Phase 1 reference  ·  Research use only", size=10, color=VIOLET, bold=True)

    doc.add_paragraph()
    metadata = [
        ("Document ID", "ONCOGON-OSIEL-SRS-001"),
        ("Version", "1.2 — Governed Professor, model/redocking qualification and interactive 3D"),
        ("Date", "22 August 2026"),
        ("Status", "Scope freeze and implementation baseline"),
        ("System boundary", "In-silico research prioritization; not clinical decision support"),
    ]
    table = doc.add_table(rows=len(metadata), cols=2)
    table.style = "Table Grid"
    set_table_widths(table, [1.6, 4.9])
    for row, (label, value) in zip(table.rows, metadata):
        set_cell_shading(row.cells[0], LIGHT)
        row.cells[0].text = label
        row.cells[1].text = value
        for run in row.cells[0].paragraphs[0].runs:
            set_font(run, size=9, color=NAVY, bold=True)
        for run in row.cells[1].paragraphs[0].runs:
            set_font(run, size=9, color="333333")
    doc.add_paragraph()
    boundary = doc.add_paragraph()
    boundary.paragraph_format.space_before = Pt(12)
    boundary.paragraph_format.space_after = Pt(0)
    add_runs(
        boundary,
        "This document distinguishes executable reference behavior from production adapters and scientifically validated capabilities. Simulated results are never efficacy evidence.",
        size=10,
        color=MUTED,
        italic=True,
    )
    doc.add_page_break()

    h = doc.add_heading("Document map", level=1)
    h.paragraph_format.space_before = Pt(0)
    sections = [
        "1–5  Executive decision, scope, roles, and use cases",
        "6  Requirements catalogue — OSIEL-REQ-001 through OSIEL-REQ-067",
        "7–13  Architecture, data, chemistry, models, ranking, and experiments",
        "14–18  APIs, events, sources, frontend, and security",
        "19–25  Non-functional requirements, deployment, verification, risks, and definition of done",
        "Appendices  Official technical references and scientific handoff checklist",
    ]
    for item in sections:
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, item, size=11)
    note = doc.add_table(rows=1, cols=1)
    set_table_widths(note, [6.5])
    set_cell_shading(note.cell(0, 0), PALE_GREEN)
    p = note.cell(0, 0).paragraphs[0]
    add_runs(p, "Implementation baseline: 10.10B provider map · 192 local references · governed Professor · ChEMBL model lab · Vina + 0.8636 Å RMSD · interactive 3D · 52 backend tests.", size=10, color=GREEN, bold=True)
    doc.add_page_break()


def clean_inline(text: str) -> str:
    return text.strip().replace("  ", " ")


def table_width_pattern(headers: list[str]) -> list[float]:
    normalized = [header.lower() for header in headers]
    if len(headers) == 4 and normalized[0] == "id":
        return [0.95, 2.55, 2.00, 1.00]
    if len(headers) == 4:
        return [1.35, 1.8, 2.25, 1.1]
    if len(headers) == 3:
        if "weight" in normalized:
            return [1.70, 0.80, 4.00]
        if "method and path" in normalized:
            return [2.05, 2.35, 2.10]
        return [1.55, 3.65, 1.3]
    if len(headers) == 2:
        return [1.8, 4.7]
    return [6.5 / len(headers)] * len(headers)


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    headers = rows[0]
    status_table = "status" in headers[-1].lower()
    table = doc.add_table(rows=len(rows), cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, table_width_pattern(headers))
    set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        prevent_row_split(table.rows[row_index])
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            if row_index == 0:
                set_cell_shading(cell, LIGHT)
                add_runs(p, clean_inline(value), size=8.5, color=NAVY, bold=True)
            else:
                fill = WHITE
                lower = value.lower()
                if status_table and col_index == len(values) - 1:
                    if "implemented" in lower:
                        fill = PALE_GREEN
                    elif "deferred" in lower:
                        fill = PALE_RED
                    elif "partial" in lower or "contract" in lower or "adapter" in lower:
                        fill = PALE_AMBER
                set_cell_shading(cell, fill)
                add_runs(p, clean_inline(value), size=8.3, color="303846", bold=col_index == 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def create_numbering_instance(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    number_format = OxmlElement("w:numFmt")
    number_format.set(qn("w:val"), "decimal")
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    paragraph_properties = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    indentation = OxmlElement("w:ind")
    indentation.set(qn("w:left"), "720")
    indentation.set(qn("w:hanging"), "360")
    paragraph_properties.extend([tabs, indentation])
    level.extend([start, number_format, level_text, justification, paragraph_properties])
    abstract.append(level)
    numbering.append(abstract)

    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(num_id))
    abstract_reference = OxmlElement("w:abstractNumId")
    abstract_reference.set(qn("w:val"), str(abstract_id))
    number.append(abstract_reference)
    numbering.append(number)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is not None:
        p_pr.remove(num_pr)
    num_pr = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    num_pr.extend([level, number])
    p_pr.insert(0, num_pr)


def parse_markdown_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    table_lines: list[str] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        table_lines.append(lines[index].strip())
        index += 1
    rows = []
    for line_index, line in enumerate(table_lines):
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        if line_index == 1 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        rows.append(cells)
    return rows, index


def add_body_from_markdown(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    start = next(index for index, line in enumerate(lines) if line.startswith("## 1."))
    index = start
    active_numbering_id: int | None = None
    compact_checklist_override = False
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if not line or line == "---":
            active_numbering_id = None
            index += 1
            continue
        if line.startswith("|"):
            active_numbering_id = None
            rows, index = parse_markdown_table(lines, index)
            add_markdown_table(doc, rows)
            continue
        heading_match = re.match(r"^(#{2,4})\s+(.+)$", line)
        if heading_match:
            active_numbering_id = None
            marks, text = heading_match.groups()
            level = min(3, len(marks) - 1)
            doc.add_heading(text, level=level)
            # Named override: the final scientific handoff checklist uses tighter
            # list rhythm so the specification does not end on an orphaned item.
            compact_checklist_override = text.startswith("Appendix B")
            index += 1
            continue
        if line.startswith(">"):
            active_numbering_id = None
            table = doc.add_table(rows=1, cols=1)
            set_table_widths(table, [6.5])
            set_cell_shading(table.cell(0, 0), PALE_VIOLET)
            p = table.cell(0, 0).paragraphs[0]
            add_runs(p, line.lstrip("> ").strip(), size=11, color=VIOLET, bold=True)
            index += 1
            continue
        if line.startswith("- "):
            active_numbering_id = None
            p = doc.add_paragraph(style="List Bullet")
            if compact_checklist_override:
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.10
            add_runs(p, line[2:], size=11)
            index += 1
            continue
        if re.match(r"^\d+\.\s", line):
            if active_numbering_id is None:
                active_numbering_id = create_numbering_instance(doc)
            p = doc.add_paragraph(style="List Number")
            apply_numbering(p, active_numbering_id)
            add_runs(p, re.sub(r"^\d+\.\s+", "", line), size=11)
            index += 1
            continue
        active_numbering_id = None
        p = doc.add_paragraph()
        add_runs(p, line, size=11)
        index += 1


def audit(doc: Document) -> None:
    section = doc.sections[0]
    assert round(section.page_width.inches, 2) == 8.5
    assert round(section.page_height.inches, 2) == 11.0
    assert all(round(value.inches, 2) == 1.0 for value in (section.top_margin, section.bottom_margin, section.left_margin, section.right_margin))
    assert doc.styles["Normal"].font.name == "Calibri"
    assert doc.styles["Normal"].font.size.pt == 11
    for table in doc.tables:
        assert len(table.rows) >= 1


def main(source: Path, output: Path) -> None:
    markdown = source.read_text(encoding="utf-8")
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_body_from_markdown(doc, markdown)
    audit(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(f"Wrote {output} with {len(doc.paragraphs)} paragraphs and {len(doc.tables)} tables")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_srs_docx.py input.md output.docx")
    main(Path(sys.argv[1]), Path(sys.argv[2]))
